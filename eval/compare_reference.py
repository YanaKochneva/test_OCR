from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document as WordDocument


def _tokens(text: str) -> list[str]:
    return re.findall(r"\w+", text.lower().replace("ё", "е"), flags=re.UNICODE)


def _reference_text(path: Path) -> str:
    reference = WordDocument(path)
    parts = [paragraph.text for paragraph in reference.paragraphs]
    for table in reference.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return " ".join(parts)


def _ir_text(data: dict[str, Any]) -> str:
    parts: list[str] = []
    for page in data.get("pages", []):
        for block in page.get("blocks", []):
            parts.append(block.get("text") or "")
            table = block.get("table") or {}
            parts.extend(cell.get("text") or "" for cell in table.get("cells", []))
    return " ".join(parts)


def compare(reference_path: Path, ir_path: Path) -> dict[str, Any]:
    reference = _reference_text(reference_path)
    ir = _ir_text(json.loads(ir_path.read_text(encoding="utf-8")))
    expected = Counter(_tokens(reference))
    actual = Counter(_tokens(ir))
    matched = sum((expected & actual).values())
    precision = matched / sum(actual.values()) if actual else 0.0
    recall = matched / sum(expected.values()) if expected else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "metric": "multiset exact-token overlap; ignores order and page alignment",
        "reference_chars": len(reference),
        "ir_chars": len(ir),
        "reference_tokens": sum(expected.values()),
        "ir_tokens": sum(actual.values()),
        "matched_tokens": matched,
        "precision": precision,
        "recall": recall,
        "f1": f1,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare OCR IR text with a DOCX reference.")
    parser.add_argument("reference", type=Path, help="Reference .docx file")
    parser.add_argument("ir", type=Path, help="Generated document.json")
    parser.add_argument("--output", type=Path, help="Optional JSON report path")
    args = parser.parse_args()

    if not args.reference.is_file():
        parser.error(f"DOCX reference does not exist: {args.reference}")
    if not args.ir.is_file():
        parser.error(f"IR JSON does not exist: {args.ir}")

    report = compare(args.reference, args.ir)
    rendered = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
