"""Sequential, isolated OCR runs against paired reference documents."""
from __future__ import annotations

import csv
import json
import subprocess
import sys
import time
import math
from pathlib import Path
from typing import Any

import pypdfium2 as pdfium
from docx import Document as WordDocument

from docpipe.evaluation import _ir_text, _reference_text, compare_text

SUPPORTED = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


def reference_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return _reference_text(path)
    if path.suffix.lower() != ".pdf":
        return ""
    parts = []
    with pdfium.PdfDocument(str(path)) as document:
        for index in range(len(document)):
            page = document[index]
            textpage = page.get_textpage()
            try:
                parts.append(textpage.get_text_range())
            finally:
                textpage.close()
                page.close()
    return "\n".join(parts)


def discover_pairs(root: Path) -> list[dict[str, Any]]:
    """Each directory containing supported files is one case, including nested ones."""
    if not root.is_dir():
        raise ValueError(f"Dataset directory does not exist: {root}")
    cases = []
    for folder in sorted(root.rglob("*")):
        if not folder.is_dir():
            continue
        files = sorted(p for p in folder.iterdir() if p.is_file()
                       and p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
        if not files:
            continue
        case: dict[str, Any] = {"case": folder.relative_to(root).as_posix()}
        try:
            if len(files) != 2:
                raise ValueError(f"Expected exactly 2 supported documents, found {len(files)}")
            texts = [reference_text(p) for p in files]
            references = [i for i, text in enumerate(texts) if text.strip()]
            if len(references) != 1:
                raise ValueError("Pair must contain exactly one text-bearing reference and one text-free scan")
            ref = references[0]
            scan = files[1 - ref]
            if scan.suffix.lower() == ".docx":
                raise ValueError("A DOCX without extractable text is not a supported scan")
            case.update(reference=str(files[ref]), scan=str(scan), status="ready")
        except Exception as exc:
            case.update(status="invalid_pair", error=str(exc))
        cases.append(case)
    if not cases:
        raise ValueError("No document pairs found in subdirectories")
    return cases


def score_structure(reference: Path, ir: dict[str, Any], output: Path) -> dict[str, Any]:
    """DOCX supplies native table structure; PDFs require annotated gold for this."""
    if reference.suffix.lower() != ".docx":
        return {"status": "not_available", "reason": "PDF table/figure gold annotations required"}
    expected = WordDocument(reference)
    actual = WordDocument(output / "document.docx")
    tables = [b["table"] for p in ir["pages"] for b in p["blocks"] if b.get("table")]
    signatures = [(len(t.rows), len(t.columns)) for t in expected.tables]
    detected = [(t["n_rows"], t["n_cols"]) for t in tables]
    matched = sum(a == b for a, b in zip(signatures, detected))
    return {
        "reference_tables": len(signatures), "ocr_tables": len(detected),
        "output_native_tables": len(actual.tables),
        "table_dimensions_accuracy": matched / max(len(signatures), len(detected))
        if signatures or detected else None,
        "reference_table_dimensions": signatures, "ocr_table_dimensions": detected,
        "reference_inline_images": len(expected.inline_shapes),
        # Counts are diagnostics, not image detection precision/recall.
        "ocr_figures": sum(b.get("type") == "figure" for p in ir["pages"] for b in p["blocks"]),
    }


def run_benchmark(root: Path, out: Path, *, engine: str = "docling",
                  config: Path | None = None, timeout: int = 1800,
                  dry_run: bool = False, selected_cases: list[str] | None = None) -> dict[str, Any]:
    root, out = root.resolve(), out.resolve()
    if out == root or out.is_relative_to(root) or root.is_relative_to(out):
        raise ValueError("Report directory and dataset must not contain one another")
    if timeout <= 0:
        raise ValueError("Timeout must be positive")
    if out.exists() and any(out.iterdir()):
        raise ValueError("Report directory must be empty: existing results will not be overwritten")
    cases = discover_pairs(root)
    if selected_cases:
        available = {case["case"] for case in cases}
        missing = set(selected_cases) - available
        if missing:
            raise ValueError(f"Unknown cases: {sorted(missing)}")
        cases = [case for case in cases if case["case"] in selected_cases]
    out.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "schema_version": 1, "dataset": str(root), "engine": engine,
        "dry_run": dry_run, "cases": cases,
        "metrics_note": "CER/WER normalize case, whitespace and yo/e; token F1 ignores order. "
        "Table dimensions compare tables in document order. No acceptance thresholds configured. "
        "IoU and figure F1 require gold object annotations; not inferred from counts.",
    }
    for index, case in enumerate(cases):
        if case["status"] != "ready" or dry_run:
            continue
        case_out = out / f"case-{index + 1:04d}"
        case_out.mkdir()
        command = [sys.executable, "-m", "docpipe.cli", "parse", case["scan"],
                   "--engine", engine, "--out", str(case_out), "--outputs", "docx", "pdf-positional"]
        if config:
            command.extend(["--config", str(config.resolve())])
        started = time.perf_counter()
        try:
            with (case_out / "run.log").open("w", encoding="utf-8") as log:
                result = subprocess.run(command, stdout=log, stderr=subprocess.STDOUT,
                                        timeout=timeout, check=False)
            case["exit_code"] = result.returncode
            ir_path = case_out / "document.json"
            if not ir_path.is_file():
                raise RuntimeError(f"OCR did not produce document.json; exit code {result.returncode}")
            ir = json.loads(ir_path.read_text(encoding="utf-8"))
            case["text"] = compare_text(reference_text(Path(case["reference"])), _ir_text(ir))
            case["structure"] = score_structure(Path(case["reference"]), ir, case_out)
            verification_path = case_out / "verification.json"
            case["verification"] = (json.loads(verification_path.read_text(encoding="utf-8"))
                                    if verification_path.exists() else None)
            case["status"] = "evaluated" if result.returncode == 0 else "processing_failed"
        except subprocess.TimeoutExpired:
            case.update(status="timeout", error=f"Processing exceeded {timeout} seconds")
        except Exception as exc:
            case.update(status="processing_failed", error=str(exc))
        finally:
            case["elapsed_seconds"] = round(time.perf_counter() - started, 3)
            case["output"] = str(case_out)
            _write_report(out, report)
    scored = [c["text"] for c in cases if "text" in c]
    report["summary"] = {
        "cases": len(cases),
        "statuses": {status: sum(c["status"] == status for c in cases)
                     for status in sorted({c["status"] for c in cases})},
        "micro_cer": sum(s["cer"] * s["reference_chars"] for s in scored)
        / sum(s["reference_chars"] for s in scored) if scored else None,
        "micro_wer": sum(s["wer"] * s["reference_tokens"] for s in scored)
        / sum(s["reference_tokens"] for s in scored) if scored else None,
        "scored_cases": len(scored),
        "macro_cer": sum(s["cer"] for s in scored) / len(scored) if scored else None,
        "macro_wer": sum(s["wer"] for s in scored) / len(scored) if scored else None,
        "worst_document_cer": max((s["cer"] for s in scored), default=None),
        "document_exact_match_rate": sum(s["document_exact_match"] for s in scored) / len(scored)
        if scored else None,
    }
    completed = [c for c in cases if c["status"] == "evaluated"]
    latencies = sorted(c["elapsed_seconds"] for c in completed)
    report["summary"]["p95_document_seconds"] = (
        latencies[math.ceil(.95 * len(latencies)) - 1] if latencies else None
    )
    report["runtime_note"] = "Sequential isolated processes; cold engine initialization and exports included. "
    report["runtime_note"] += "p95 uses nearest-rank over successful runs only; failures/timeouts reported separately."
    _write_report(out, report)
    return report


def _write_report(out: Path, report: dict[str, Any]) -> None:
    temporary = out / "report.json.tmp"
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(out / "report.json")
    with (out / "summary.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        fields = ["case", "status", "cer", "wer", "strict_cer", "strict_wer", "f1",
                  "document_exact_match", "elapsed_seconds", "error"]
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for case in report["cases"]:
            writer.writerow({key: case.get("text", {}).get(key, case.get(key)) for key in fields})
