from __future__ import annotations

from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt,RGBColor
from docx.text.paragraph import Paragraph

from docpipe.ir import Block, BlockType, Document, Page
from docpipe.renderers.table_layout import build_table_layout
from docpipe.renderers.text_layout import wrap_text
from docpipe.renderers.typography import (
    BODY_FONT_PT,
    BODY_LINE_SPACING,
    FONT_FAMILY,
    TABLE_FONT_PT,
    text_style,
)

_EMU_PER_POINT = 12_700
_TWIPS_PER_POINT = 20
_TABLE_BORDER_PT = .5


def _set_page(section, page: Page) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if page.width_pt>page.height_pt else WD_ORIENT.PORTRAIT
    section.page_width = Pt(page.width_pt)
    section.page_height = Pt(page.height_pt)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    grid = section._sectPr.find(qn('w:docGrid'))
    if grid is not None:
        section._sectPr.remove(grid)


def _anchor_paragraph(doc: WordDocument):
    existing = getattr(doc, '_docpipe_page_anchor', None)
    if existing is not None:
        return existing
    paragraph = doc.add_paragraph()
    if 'DocpipeAnchor' in doc.styles:
        paragraph.style = doc.styles['DocpipeAnchor']
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = False
    mark = OxmlElement('w:rPr')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '2')
    mark.append(size)
    paragraph._p.get_or_add_pPr().append(mark)
    # One stable logical page anchor, rather than one body paragraph per object.
    doc._docpipe_page_anchor = paragraph
    return paragraph


def _add_textbox(
    doc: WordDocument,
    *,
    shape_id: str,
    x: float,
    y: float,
    width: float,
    height: float,
    z_index: int,
    text: str,
    font_size: float,
    bold: bool = False,
    italic: bool = False,
    border: bool = False,
    table_content=None,
    font_family: str = FONT_FAMILY,
    horizontal_scale: int | None = None,
    source_line=None,
) -> None:
    lines = text.splitlines() or [""]
    escaped_lines = []
    for line_index, line in enumerate(lines):
        if line_index:
            escaped_lines.append("<w:br/>")
        escaped_lines.append(f'<w:t xml:space="preserve">{escape(line)}</w:t>')
    bold_xml = "<w:b/>" if bold else ""
    italic_xml = "<w:i/>" if italic else ""
    font_half_points = max(2, round(font_size * 2))
    scale_xml=f'<w:w w:val="{horizontal_scale}"/>' if horizontal_scale else ''
    namespaces = nsdecls('w', 'wp', 'a') + ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    cx, cy = round(max(width,.5)*_EMU_PER_POINT), round(max(height,.5)*_EMU_PER_POINT)
    line = '<a:solidFill><a:srgbClr val="333333"/></a:solidFill>' if border else '<a:noFill/>'
    xml = f"""<w:drawing {namespaces}>
      <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{max(0,z_index)}" behindDoc="0" locked="0" layoutInCell="0" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page"><wp:posOffset>{round(x*_EMU_PER_POINT)}</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="page"><wp:posOffset>{round(y*_EMU_PER_POINT)}</wp:posOffset></wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/><wp:docPr id="{doc.part.next_id}" name="{escape(shape_id)}"/>
        <wp:cNvGraphicFramePr/>
        <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
          <wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>
            <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
            <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="6350">{line}</a:ln>
          </wps:spPr><wps:txbx><w:txbxContent>
          <w:p><w:pPr><w:snapToGrid w:val="0"/><w:spacing w:before="0" w:after="0" w:line="{max(20, round(font_half_points * BODY_LINE_SPACING * 10))}" w:lineRule="exact"/></w:pPr>
            <w:r><w:rPr><w:rFonts w:ascii="{escape(font_family)}" w:hAnsi="{escape(font_family)}" w:eastAsia="{escape(font_family)}" w:cs="{escape(font_family)}"/>{bold_xml}{italic_xml}{scale_xml}<w:sz w:val="{font_half_points}"/></w:rPr>{''.join(escaped_lines)}</w:r>
          </w:p>
          </w:txbxContent></wps:txbx>
          <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="12700" rIns="12700" tIns="0" bIns="0" anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr>
          </wps:wsp>
        </a:graphicData></a:graphic>
      </wp:anchor>
    </w:drawing>"""
    drawing = parse_xml(xml)
    if source_line is not None:
        body=drawing.find('.//{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}bodyPr')
        body.set('lIns','0');body.set('rIns','0')
        body.set('wrap','none')
    if source_line is not None and source_line.rotation in (90,270):
        body=drawing.find('.//{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}bodyPr')
        body.set('vert','vert270' if source_line.rotation==90 else 'vert')
    if source_line is not None:
        node=drawing.find('.//' + qn('w:txbxContent')).find(qn('w:p'))
        for child in list(node):
            if child.tag!=qn('w:pPr'):
                node.remove(child)
        _append_line_runs(Paragraph(node,doc),source_line,font_size)
    if table_content is not None:
        container = drawing.find('.//' + qn('w:txbxContent'))
        container.clear()
        container.append(table_content)
        # A table in a text frame requires a trailing, minimal paragraph.
        container.append(parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="DocpipeAnchor"/><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/><w:rPr><w:sz w:val="2"/></w:rPr></w:pPr><w:r><w:rPr><w:sz w:val="2"/></w:rPr><w:t xml:space="preserve"> </w:t></w:r></w:p>'))
        body = drawing.find('.//{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}bodyPr')
        body.set('lIns','0'); body.set('rIns','0')
    run = OxmlElement("w:r")
    run.append(drawing)
    paragraph = _anchor_paragraph(doc)
    paragraph._p.append(run)


def _float_picture(
    picture,
    x: float,
    y: float,
    width: float,
    height: float,
    z_index: int,
) -> None:
    inline = picture._inline
    drawing = inline.getparent()
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": str(max(0, z_index)),
        "behindDoc": "0", "locked": "0", "layoutInCell": "0", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)

    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    offset_h = OxmlElement("wp:posOffset")
    offset_h.text = str(round(x * _EMU_PER_POINT))
    position_h.append(offset_h)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = str(round(y * _EMU_PER_POINT))
    position_v.append(offset_v)
    extent = inline.find(qn("wp:extent"))
    if extent is not None:
        extent.set("cx", str(round(width * _EMU_PER_POINT)))
        extent.set("cy", str(round(height * _EMU_PER_POINT)))
    effect_extent = inline.find(qn("wp:effectExtent"))
    wrap = OxmlElement("wp:wrapNone")
    doc_pr = inline.find(qn("wp:docPr"))
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    graphic_extents = graphic.findall(".//" + qn("a:ext")) if graphic is not None else []
    if graphic_extents:
        graphic_extents[0].set("cx", str(round(width * _EMU_PER_POINT)))
        graphic_extents[0].set("cy", str(round(height * _EMU_PER_POINT)))
    ordered = [simple, position_h, position_v, extent]
    if effect_extent is not None:
        ordered.append(effect_extent)
    ordered.extend([wrap, doc_pr, frame, graphic])
    for element in ordered:
        if element is not None:
            if element.getparent() is inline:
                inline.remove(element)
            anchor.append(element)
    drawing.replace(inline, anchor)


def _set_table_borders(table, color='333333', width_pt=_TABLE_BORDER_PT) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        layout = tbl_pr.find(qn("w:tblLayout"))
        index = tbl_pr.index(layout) if layout is not None else len(tbl_pr)
        tbl_pr.insert(index, borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(max(1,round(width_pt*8))))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _anchor_table(table, block: Block, top: float | None = None) -> None:
    """Float a native Word table at the OCR table's page coordinates."""
    box = block.bbox
    tbl_pr = table._tbl.tblPr
    position = tbl_pr.find(qn("w:tblpPr"))
    if position is None:
        position = OxmlElement("w:tblpPr")
        style = tbl_pr.find(qn("w:tblStyle"))
        tbl_pr.insert(tbl_pr.index(style) + 1 if style is not None else 0, position)
    for name, value in {
        "horzAnchor": "page",
        "vertAnchor": "page",
        "tblpX": str(round(box.x0 * _TWIPS_PER_POINT)),
        "tblpY": str(round((box.y0 if top is None else top) * _TWIPS_PER_POINT)),
        "leftFromText": "0",
        "rightFromText": "0",
        "topFromText": "0",
        "bottomFromText": "0",
    }.items():
        position.set(qn(f"w:{name}"), value)
    overlap = tbl_pr.find(qn('w:tblOverlap'))
    if overlap is None:
        overlap = OxmlElement('w:tblOverlap')
        tbl_pr.insert(tbl_pr.index(position)+1, overlap)
    overlap.set(qn('w:val'), 'overlap')


def _append_line_runs(paragraph,line,fallback_size):
    from docpipe.renderers.fonts import font_path,require_font
    from PIL import ImageFont
    words=line.runs or [line]
    comment_runs=[]
    for index,word in enumerate(words):
        size=word.font_size_pt or line.font_size_pt or fallback_size
        family=word.font_family or line.font_family or FONT_FAMILY
        path=font_path(family,word.bold,word.italic) or require_font()
        font=ImageFont.truetype(str(path),100)
        glyph=font.getbbox(word.text,anchor='ls')
        width=(glyph[2]-glyph[0])*size/100
        available=word.bbox.y1-word.bbox.y0 if word.rotation in (90,270) else word.bbox.x1-word.bbox.x0
        ratio=word.horizontal_scale_pct/100 if word.horizontal_scale_pct else available/width if width else 1
        run=paragraph.add_run(word.text)
        comment_runs.append(run)
        run.font.name=family;run.font.size=Pt(size);run.bold=word.bold;run.italic=word.italic
        if word.color or line.color:
            run.font.color.rgb=RGBColor.from_string((word.color or line.color)[1:])
        scale=OxmlElement('w:w');scale.set(qn('w:val'),str(max(1,round(ratio*100))))
        run._r.get_or_add_rPr().append(scale)
        if index+1<len(words):
            space=paragraph.add_run(' ');space.font.name=family;space.font.size=Pt(size)
            space.bold=word.bold;space.italic=word.italic
            desired=words[index+1].bbox.x0-word.bbox.x1
            tail=(font.getlength(word.text)-glyph[2])*size/100*ratio
            tracking=OxmlElement('w:spacing')
            tracking.set(qn('w:val'),str(round((desired-tail-font.getlength(' ')*size/100)*20)))
            space._r.get_or_add_rPr().append(tracking)
    from docpipe.style_review import line_review
    reason=line_review(line)
    if reason and comment_runs:
        paragraph.part.document.add_comment(comment_runs,'Проверить оформление: '+reason,
                                            author='DocPipe',initials='OCR')


def _format_table_cell(cell, text: str, font_size: float = TABLE_FONT_PT, source=None) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        alignment = cell_properties.find(qn('w:vAlign'))
        cell_properties.insert(cell_properties.index(alignment) if alignment is not None
                               else len(cell_properties), shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), (source.background_color or '#FFFFFF')[1:] if source else 'FFFFFF')
    if source and source.lines:
        cell.text=''
        for index,line in enumerate(source.lines):
            paragraph=cell.paragraphs[0] if index==0 else cell.add_paragraph()
            size=line.font_size_pt or font_size
            paragraph.paragraph_format.space_before=Pt(max(0,line.bbox.y0-source.bbox.y0-1-size*.2) if index==0 else 0)
            paragraph.paragraph_format.space_after=Pt(0)
            paragraph.paragraph_format.left_indent=Pt(max(0,line.bbox.x0-source.bbox.x0-2))
            paragraph.paragraph_format.line_spacing=Pt(
                source.lines[index+1].bbox.y0-line.bbox.y0 if index+1<len(source.lines) else size)
            paragraph.paragraph_format.widow_control=False
            paragraph.paragraph_format.keep_together=False
            _append_line_runs(paragraph,line,size)
        return
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(font_size * BODY_LINE_SPACING)
        paragraph.paragraph_format.widow_control = False
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'),'0')
        props=paragraph._p.get_or_add_pPr()
        spacing=props.find(qn('w:spacing'))
        props.insert(props.index(spacing) if spacing is not None else len(props),snap)
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(font_size)


def _word_table_layout(block: Block):
    source_height = block.bbox.y1-block.bbox.y0
    for half_points in range(round(TABLE_FONT_PT*2),7,-1):
        size = half_points/2
        layout = build_table_layout(block,font_size_pt=size,full_line_boxes=True)
        if sum(layout.row_heights) <= source_height+1:
            return build_table_layout(block,font_size_pt=size,preserve_geometry=True),size
    return build_table_layout(block,font_size_pt=4,preserve_geometry=True),4.0


def _add_table(doc: WordDocument, block: Block, top: float | None = None, page_index: int = 0) -> None:
    layout,font_size = _word_table_layout(block)
    table = doc.add_table(rows=layout.rows, cols=layout.cols)
    table.autofit = False
    # Native table is placed inside its page-anchored frame, not in body flow.
    border_pt=round((block.table.border_width_pt or _TABLE_BORDER_PT)*8)/8
    _set_table_borders(table,(block.table.border_color or '#333333')[1:],border_pt)

    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is not None:
        tbl_width.set(qn("w:w"), str(round((block.bbox.x1 - block.bbox.x0) * _TWIPS_PER_POINT)))
        tbl_width.set(qn("w:type"), "dxa")

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        look = tbl_pr.find(qn("w:tblLook"))
        tbl_pr.insert(tbl_pr.index(look) if look is not None else len(tbl_pr), margins)
    for edge, value in (("top", 20), ("bottom", 20), ("start", 40), ("end", 40)):
        item = margins.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")

    for col_index, width in enumerate(layout.column_widths):
        table.columns[col_index].width = Pt(width)
        for row in table.rows:
            row.cells[col_index].width = Pt(width)
    for row_index, height in enumerate(layout.row_heights):
        row = table.rows[row_index]
        # Native row layout includes both horizontal cell borders. Compensate
        # for their thickness instead of accumulating 1pt drift per source row.
        row.height = Pt(max(1,height-2*border_pt))
        # Word must not enlarge a measured row and clip later rows in its frame.
        # Layout reserved complete line boxes and cell padding before this point.
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            _format_table_cell(cell, "", font_size)

    for placement in layout.cells:
        cell = table.cell(placement.row, placement.col)
        if placement.rowspan > 1 or placement.colspan > 1:
            last = table.cell(
                placement.row + placement.rowspan - 1,
                placement.col + placement.colspan - 1,
            )
            cell = cell.merge(last)
        source=next((c for c in block.table.cells if (c.row,c.col)==(placement.row,placement.col)),None)
        _format_table_cell(cell, placement.text, font_size, source)
    _add_textbox(doc,shape_id=f'docpipe-table-{page_index}-{block.order}',
        x=block.bbox.x0,y=block.bbox.y0 if top is None else top,
        width=block.bbox.x1-block.bbox.x0,
        height=sum(layout.row_heights)+2,z_index=100+block.order,text='',
        font_size=1,table_content=table._tbl)


def _display_text(block: Block) -> str:
    text = block.text or ''
    if block.type==BlockType.LIST_ITEM and not re.match(r'^\s*(?:[•●▪◦]|[-–—]\s|\d+[.)]\s)',text):
        text='• '+text
    return text


def _word_style(block: Block):
    size,bold,italic=text_style(block.type.value)
    if block.type in {BlockType.HEADING,BlockType.CAPTION} and block.text:
        width=max(1,block.bbox.x1-block.bbox.x0)
        expected=max(1,round((block.bbox.y1-block.bbox.y0)/size))
        while size>8 and len(wrap_text(_display_text(block),width,(size,bold,italic)).splitlines())>expected:
            size-=.5
    width = max(1, block.bbox.x1-block.bbox.x0)
    height = max(1, block.bbox.y1-block.bbox.y0)
    while size > 4 and len(wrap_text(block.text or '',width,(size,bold,italic)).splitlines())*size+2 > height:
        size -= .5
    return size,bold,italic


def _word_text_height(block: Block) -> float:
    if block.table:
        return sum(_word_table_layout(block)[0].row_heights) + 2
    if not block.text:
        return block.bbox.y1-block.bbox.y0
    text = _display_text(block)
    style = _word_style(block)
    lines = max(1,len(wrap_text(text,block.bbox.x1-block.bbox.x0,style).splitlines()))
    # Word requires complete line boxes, not just the ink/glyph bounding height.
    return max(block.bbox.y1-block.bbox.y0, lines*style[0]*BODY_LINE_SPACING + 4)


def _add_block(doc: WordDocument, block: Block, page: Page, base: Path, text_top: float | None = None) -> None:
    box = block.bbox
    width = max(0.5, box.x1 - box.x0)
    height = max(0.5, box.y1 - box.y0)
    z_index = 100 + block.order
    if block.type == BlockType.FIGURE and block.figure:
        path = base / block.figure.file
        if not path.is_file():
            raise FileNotFoundError(f"Figure image not found: {path}")
        paragraph = _anchor_paragraph(doc)
        run = paragraph.add_run()
        run.font.size = Pt(1)
        picture = run.add_picture(str(path), width=Pt(width), height=Pt(height))
        _float_picture(picture, box.x0, box.y0, width, height, z_index)
        return
    if block.type == BlockType.TABLE and block.table:
        _add_table(doc, block, text_top, page.index)
        return
    if block.text:
        if block.lines:
            for index, line in enumerate(block.lines):
                local = block.model_copy(update={'bbox':line.bbox,'text':line.text,'lines':None})
                size, bold, italic = _word_style(local)
                size=line.font_size_pt or size
                family=line.font_family or FONT_FAMILY
                if line.font_family:
                    bold,italic=line.bold,line.italic
                from docpipe.renderers.fonts import font_path,require_font
                from PIL import ImageFont
                path=font_path(family,bold,italic) or require_font()
                measured=ImageFont.truetype(str(path),100).getlength(line.text)*size/100
                scale=round(line.horizontal_scale_pct) if line.horizontal_scale_pct is not None else (max(1,round((line.bbox.x1-line.bbox.x0)/measured*100)) if measured else None)
                _add_textbox(doc,shape_id=f'docpipe-line-{page.index}-{block.order}-{index}',
                    x=line.bbox.x0,y=line.bbox.y0,
                    width=max(.5,line.bbox.x1-line.bbox.x0),
                    height=max(size*1.2,line.bbox.y1-line.bbox.y0),z_index=z_index,
                    text=line.text,font_size=size,bold=bold,italic=italic,
                    font_family=family,horizontal_scale=scale,source_line=line)
            return
        text = _display_text(block)
        font_size, bold, italic = _word_style(block)
        text = wrap_text(text, width, (font_size, bold, italic))
        content_height = height
        _add_textbox(
            doc,
            shape_id=f"docpipe-text-{page.index}-{block.order}",
            x=box.x0,
            y=box.y0 if text_top is None else text_top,
            width=width,
            height=content_height,
            z_index=z_index,
            text=text,
            font_size=font_size,
            bold=bold,
            italic=italic,
        )


def render(document: Document, out: Path, base_dir: Path | None = None) -> Path:
    """Create a page-faithful DOCX with objects anchored at their IR coordinates."""
    out.parent.mkdir(parents=True, exist_ok=True)
    from docpipe.renderers.fixed_layout import stable_document
    document = stable_document(document)
    base = base_dir or out.parent
    doc = WordDocument()
    doc.styles["Normal"].font.name = FONT_FAMILY
    doc.styles["Normal"].font.size = Pt(BODY_FONT_PT)
    anchor_style=doc.styles.add_style('DocpipeAnchor',WD_STYLE_TYPE.PARAGRAPH)
    anchor_style.font.name=FONT_FAMILY
    anchor_style.font.size=Pt(1)
    anchor_style.paragraph_format.space_before=Pt(0)
    anchor_style.paragraph_format.space_after=Pt(0)
    anchor_style.paragraph_format.line_spacing=Pt(1)
    if document.pages:
        _set_page(doc.sections[0], document.pages[0])
    for index, page in enumerate(document.pages):
        doc._docpipe_page_anchor = None
        _anchor_paragraph(doc)
        for block in sorted(page.blocks, key=lambda item: item.order):
            _add_block(doc, block, page, base)
        if index + 1 < len(document.pages):
            next_section = doc.add_section(WD_SECTION.NEW_PAGE)
            _set_page(next_section, document.pages[index + 1])
    doc.save(str(out))
    return out
