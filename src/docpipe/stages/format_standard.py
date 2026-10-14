"""Apply an external DOCX typography standard without copying its content."""
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import ImageFont

from docpipe.renderers.fonts import font_path, require_font


def _effective(run, paragraph, attr):
    value = getattr(run.font, attr)
    if value is not None:
        return value
    for start in (run.style, paragraph.style):
        visited = set()
        while start is not None and start.style_id not in visited:
            visited.add(start.style_id)
            value = getattr(start.font, attr)
            if value is not None:
                return value
            start = start.base_style
    return None


def extract_standard(reference: Path) -> dict:
    """Count effective role typography; reference text never enters the profile."""
    doc = Document(reference)
    counts = defaultdict(Counter)

    def collect(paragraph, role):
        outline = paragraph.style.element.find('.//' + qn('w:outlineLvl'))
        if role == 'text' and (outline is not None or paragraph.style.name.lower().startswith('heading')):
            role = 'heading'
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = _effective(run, paragraph, 'size')
            name = _effective(run, paragraph, 'name')
            if name is None or size is None:
                continue
            key = (name, size.pt, bool(_effective(run, paragraph, 'bold')),
                   bool(_effective(run, paragraph, 'italic')))
            counts[role][key] += len(run.text)

    for p in doc.paragraphs:
        collect(p, 'text')
    def tables(items, role='table'):
        for table in items:
            seen = set()
            for row in table.rows:
                header = row._tr.find('.//' + qn('w:tblHeader')) is not None
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    fill=cell._tc.find('.//' + qn('w:shd'))
                    colored=fill is not None and (fill.get(qn('w:fill')) or '').upper() not in ('','AUTO','FFFFFF')
                    for p in cell.paragraphs:
                        collect(p, 'table_header' if role=='table' and (header or colored) else role)
                    tables(cell.tables,role)
    tables(doc.tables)
    for section in doc.sections:
        for part, role in ((section.header, 'header'), (section.footer, 'footer')):
            for p in part.paragraphs:
                collect(p, role)
            tables(part.tables,role)
    if not counts['text']:
        raise ValueError('Reference has no resolvable body typography')
    result = {}
    for role, counter in counts.items():
        if counter:
            family, size, bold, italic = counter.most_common(1)[0][0]
            result[role] = dict(font_family=family, font_size_pt=size, bold=bold, italic=italic)
            # Multiple sizes are legitimate components of a standard, not OCR font guesses.
            alternatives=[dict(font_family=f,font_size_pt=s,bold=b,italic=i)
                for (f,s,b,i),weight in counter.items()
                if f==family and b==bold and i==italic and weight>=sum(counter.values())*.05]
            result[role]['sizes_pt']=sorted({x['font_size_pt'] for x in alternatives} | {size})
    return result


def apply_standard(blocks, profile):
    """Keep every box and character; flag typography that cannot fit a source box."""
    def lines(items, role):
        standard = profile.get(role, profile['text'])
        result = []
        for line in items or []:
            style={k:v for k,v in standard.items() if k!='sizes_pt'}
            path = font_path(style['font_family'], style['bold'], style['italic'])
            missing = path is None
            font = ImageFont.truetype(str(path or require_font()), 200)
            width = font.getlength(line.text) * style['font_size_pt'] / 200
            box_width = (line.bbox.y1-line.bbox.y0 if line.rotation in (90,270)
                         else line.bbox.x1-line.bbox.x0)
            for size in sorted(standard.get('sizes_pt',[]),reverse=True):
                if size<=standard['font_size_pt'] and font.getlength(line.text)*size/200<=box_width+2:
                    style['font_size_pt']=size
                    width=font.getlength(line.text)*size/200
                    break
            overflow = width > box_width + max(2, box_width * .03)
            reason = ('standard_font_unavailable' if missing else
                      'standard_text_exceeds_source_box' if overflow else None)
            restored_runs=None
            if line.runs and ' '.join(' '.join(r.text.split()) for r in line.runs)==' '.join(line.text.split()):
                # Word coordinates preserve justification without stretching the glyphs.
                restored_runs=[r.model_copy(update={**style,'runs':None,
                    'horizontal_scale_pct':100,'font_match_error':None,'font_match_margin':None,
                    'review_required':False,'review_reason':None}) for r in line.runs]
            result.append(line.model_copy(update={**style, 'runs':restored_runs,
                'font_match_error':None, 'font_match_margin':None,
                'baseline_from_top_pt':None, 'glyph_offset_x_pt':0,
                'horizontal_scale_pct':100, 'review_required':bool(reason),
                'review_reason':reason}))
        return result or None
    result=[]
    for block in blocks:
        update={'lines':lines(block.lines, block.type.value)}
        if block.table:
            update['table']=block.table.model_copy(update={'cells':[
                cell.model_copy(update={'lines':lines(cell.lines,'table_header' if cell.row==0 and
                    cell.background_color and cell.background_color.upper()!='#FFFFFF' else 'table')})
                for cell in block.table.cells]})
        result.append(block.model_copy(update=update))
    return result
