import json
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from docpipe.benchmark import discover_pairs, run_benchmark
from docpipe.evaluation import compare_text


def pair(root: Path, name="case"):
    folder = root / name
    folder.mkdir(parents=True)
    word = Document()
    word.add_paragraph("Reference text")
    word.save(folder / "reference.docx")
    pdf = canvas.Canvas(str(folder / "scan.pdf"))
    pdf.rect(10, 10, 100, 100)
    pdf.showPage()
    pdf.save()
    return folder


def test_discovery_and_dry_run(tmp_path):
    root = tmp_path / "dataset"
    pair(root, "nested/case")
    report = run_benchmark(root, tmp_path / "report", dry_run=True)
    assert report["cases"][0]["scan"].endswith("scan.pdf")
    assert report["summary"]["statuses"] == {"ready": 1}
    assert report["summary"]["micro_cer"] is None
    assert (tmp_path / "report/summary.csv").read_bytes().startswith(b"\xef\xbb\xbf")


def test_ambiguous_pair_is_not_guessed(tmp_path):
    folder = pair(tmp_path / "dataset")
    pdf = canvas.Canvas(str(folder / "scan.pdf"))
    pdf.drawString(10, 10, "Also has text")
    pdf.save()
    assert discover_pairs(tmp_path / "dataset")[0]["status"] == "invalid_pair"


def test_pdf_reference_is_detected_by_content(tmp_path):
    folder = pair(tmp_path / "dataset")
    (folder / "reference.docx").unlink()
    pdf = canvas.Canvas(str(folder / "gold.pdf"))
    pdf.drawString(10, 10, "Reference text")
    pdf.save()
    assert discover_pairs(tmp_path / "dataset")[0]["reference"].endswith("gold.pdf")


def test_empty_reference_is_not_scored_perfect():
    with pytest.raises(ValueError):
        compare_text("", "wrong")
    assert compare_text("one two", "two one")["f1"] == 1
    assert compare_text("one two", "two one")["wer"] > 0


def test_batch_continues_after_timeout_and_scores_next_case(tmp_path, monkeypatch):
    root = tmp_path / "dataset"
    pair(root, "a")
    pair(root, "b")
    calls = []

    def run(command, **kwargs):
        calls.append(command)
        if len(calls) == 1:
            raise subprocess.TimeoutExpired(command, 1)
        out = Path(command[command.index("--out") + 1])
        (out / "document.json").write_text(json.dumps({
            "pages": [{"blocks": [{"type": "text", "text": "Reference text"}]}]
        }))
        Document().save(out / "document.docx")
        (out / "verification.json").write_text('{"ok": true, "result": "PASS"}')
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr("docpipe.benchmark.subprocess.run", run)
    report = run_benchmark(root, tmp_path / "report", timeout=1)
    assert [c["status"] for c in report["cases"]] == ["timeout", "evaluated"]
    assert report["summary"]["micro_cer"] == 0
    assert report["summary"]["scored_cases"] == 1
    assert json.loads((tmp_path / "report/report.json").read_text())["summary"] == report["summary"]


def test_output_cannot_pollute_dataset_or_overwrite_results(tmp_path):
    root = tmp_path / "dataset"
    pair(root)
    with pytest.raises(ValueError):
        run_benchmark(root, root / "output", dry_run=True)
    out = tmp_path / "report"
    out.mkdir()
    (out / "existing.txt").write_text("keep")
    with pytest.raises(ValueError):
        run_benchmark(root, out, dry_run=True)
    assert (out / "existing.txt").read_text() == "keep"


def test_only_selected_cases_are_included(tmp_path):
    root = tmp_path / "dataset"
    pair(root, "doc_img_001")
    pair(root, "doc_img_003")
    pair(root, "doc_img_004")
    report = run_benchmark(root, tmp_path / "selected", dry_run=True,
                           selected_cases=["doc_img_003", "doc_img_004"])
    assert [case["case"] for case in report["cases"]] == ["doc_img_003", "doc_img_004"]
    with pytest.raises(ValueError, match="Unknown cases"):
        run_benchmark(root, tmp_path / "missing", dry_run=True, selected_cases=["unknown"])
