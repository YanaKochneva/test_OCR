from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document as WordDocument
from PIL import Image

from docpipe.ir import (
    BBox,
    Block,
    BlockType,
    Document,
    EngineInfo,
    FigureData,
    Page,
    SourceInfo,
    SourceType,
    TableCell,
    TableData,
    TextLayerStatus,
    PageStatus,
)
from docpipe.renderers.docx import render as render_docx
from docpipe.renderers.html import render as render_html


def sample_document() -> Document:
    return Document(
        source=SourceInfo(file="scan.pdf", sha256="a" * 64, type=SourceType.PDF),
        engine=EngineInfo(name="docling", version="test"),
        pages=[
            Page(
                index=0,
                width_pt=200,
                height_pt=300,
                raster_dpi=300,
                size_px=(833, 1250),
                status=PageStatus.OK,
                text_layer=TextLayerStatus.ABSENT,
                blocks=[
                    Block(
                        id="title",
                        type=BlockType.HEADING,
                        bbox=BBox(x0=20, y0=15, x1=180, y1=35),
                        order=0,
                        text="Page one",
                    ),
                    Block(
                        id="table",
                        type=BlockType.TABLE,
                        bbox=BBox(x0=20, y0=60, x1=120, y1=100),
                        order=1,
                        table=TableData(
                            n_rows=1,
                            n_cols=1,
                            cells=[
                                TableCell(
                                    row=0,
                                    col=0,
                                    text="Cell text",
                                    bbox=BBox(x0=20, y0=60, x1=120, y1=100),
                                )
                            ],
                        ),
                    ),
                    Block(
                        id="figure",
                        type=BlockType.FIGURE,
                        bbox=BBox(x0=130, y0=150, x1=180, y1=180),
                        order=2,
                        figure=FigureData(
                            file="images/sample.png",
                            format="png",
                            width_px=50,
                            height_px=30,
                        ),
                    ),
                ],
            ),
            Page(
                index=1,
                width_pt=200,
                height_pt=300,
                raster_dpi=300,
                size_px=(833, 1250),
                status=PageStatus.OK,
                text_layer=TextLayerStatus.ABSENT,
                blocks=[
                    Block(
                        id="second-title",
                        type=BlockType.HEADING,
                        bbox=BBox(x0=20, y0=15, x1=180, y1=35),
                        order=0,
                        text="Page two",
                    )
                ],
            ),
        ],
    )


def test_html_uses_page_coordinates_and_embeds_images(tmp_path: Path):
    image = tmp_path / "images" / "sample.png"
    image.parent.mkdir()
    Image.new("RGB", (50, 30), "red").save(image)
    output = render_html(sample_document(), tmp_path / "document.html")
    html = output.read_text(encoding="utf-8")

    assert html.count('class="page"') == 2
    assert 'left:20.000pt;top:60.000pt' in html
    assert 'data-file="images/sample.png"' in html
    assert "src=\"data:image/png;base64," in html
    assert "Cell text" in html


def test_docx_anchors_objects_and_embeds_images(tmp_path: Path):
    image = tmp_path / "images" / "sample.png"
    image.parent.mkdir()
    Image.new("RGB", (50, 30), "red").save(image)
    output = render_docx(sample_document(), tmp_path / "document.docx", tmp_path)
    WordDocument(output)

    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        image_members = [name for name in archive.namelist() if name.startswith("word/media/")]

    assert "<w:tblpPr" in document_xml
    assert 'w:tblpX="400"' in document_xml
    assert 'w:tblpY="1200"' in document_xml
    assert "<wp:anchor" in document_xml
    assert "1651000" in document_xml and "1905000" in document_xml
    assert "<w:txbxContent>" in document_xml
    assert "Cell text" in document_xml
    assert len(image_members) == 1


def test_pdf_multiline_uses_single_line_advance():
    from unittest.mock import Mock
    from docpipe.renderers.pdf_positional import _draw_lines
    pdf = Mock()
    _draw_lines(pdf, "one\ntwo\nthree", 10, 100, 150, 50, 11, "Helvetica")
    text = pdf.beginText.return_value
    text.setLeading.assert_called_once_with(11)
    assert text.textLine.call_count == 3
    text.moveCursor.assert_not_called()


def test_table_growth_reserves_space_for_following_text(tmp_path):
    from docpipe.renderers.table_layout import build_table_layout
    from docpipe.renderers.text_layout import collision_free_text_tops
    from docpipe.renderers.typography import text_style

    document = sample_document()
    table = document.pages[0].blocks[1]
    table.table.cells[0].text = "Long cell content " * 30
    following = Block(
        id="following", type=BlockType.TEXT,
        bbox=BBox(x0=20, y0=105, x1=120, y1=125),
        order=3, text="Following paragraph",
    )
    layout = build_table_layout(table)
    assert sum(layout.row_heights) > 40
    tops = collision_free_text_tops([table, following], lambda b: text_style(b.type.value))
    assert tops[id(following)] >= tops[id(table)] + sum(layout.row_heights) + 1.5


def test_searchable_overlay_contains_native_table_text():
    import pypdfium2 as pdfium
    from docpipe.renderers.pdf_searchable import _make_overlay
    with pdfium.PdfDocument(_make_overlay(sample_document(), 0, 200, 300)) as pdf:
        page = pdf[0]
        textpage = page.get_textpage()
        try:
            assert "Cell text" in textpage.get_text_range()
            assert "Page one" in textpage.get_text_range()
        finally:
            textpage.close()
            page.close()
