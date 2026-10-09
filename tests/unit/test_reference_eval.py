import json

from docx import Document as WordDocument
from docpipe.evaluation import compare_docx_reference
from docpipe.evaluation import compare_text


def test_strict_scores_and_error_decomposition():
    report = compare_text("Hello world", "hello world extra")
    assert report["strict_cer"] > report["cer"]
    assert report["word_errors"]["insertions"] == 1
    assert report["document_exact_match"] is False
    errors = report["character_errors"]
    assert sum(errors.values()) / report["reference_chars"] == report["cer"]


def test_reference_includes_word_content_control(tmp_path):
    from docx.oxml import OxmlElement
    from docpipe.evaluation import _reference_text
    word = WordDocument()
    word.add_paragraph("First")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Contents entry"
    run.append(text); paragraph.append(run); content.append(paragraph); control.append(content)
    word.element.body.insert(1, control)
    word.add_paragraph("Last")
    path = tmp_path / "control.docx"
    word.save(path)
    assert _reference_text(path) == "First Contents entry Last"


def test_reference_comparison_reports_multiset_overlap(tmp_path):
    reference = tmp_path / "reference.docx"
    word = WordDocument()
    word.add_paragraph("Hello world")
    word.save(reference)
    ir = tmp_path / "document.json"
    ir.write_text(json.dumps({"pages": [{"blocks": [{"text": "hello world"}]}]}), encoding="utf-8")

    report = compare_docx_reference(reference, ir)

    assert report["matched_tokens"] == 2
    assert report["precision"] == 1.0
    assert report["recall"] == 1.0
    assert report["f1"] == 1.0
    assert report["cer"] == 0.0
    assert report["wer"] == 0.0


def test_merged_cells_and_table_summary_are_not_counted_twice(tmp_path):
    reference = tmp_path / "merged.docx"
    word = WordDocument()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged text"
    word.save(reference)
    ir = tmp_path / "document.json"
    ir.write_text(json.dumps({"pages": [{"blocks": [{
        "text": "Merged text", "table": {"cells": [{"text": "Merged text"}]}
    }]}]}), encoding="utf-8")
    assert compare_docx_reference(reference, ir)["cer"] == 0.0


def test_reference_comparison_preserves_paragraph_table_order(tmp_path):
    reference = tmp_path / "ordered.docx"
    word = WordDocument()
    word.add_paragraph("First paragraph")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Middle table"
    word.add_paragraph("Last paragraph")
    word.save(reference)

    ir = tmp_path / "ordered.json"
    ir.write_text(
        json.dumps(
            {
                "pages": [
                    {
                        "blocks": [
                            {"text": "First paragraph"},
                            {"text": None, "table": {"cells": [{"text": "Middle table"}]}},
                            {"text": "Last paragraph"},
                        ]
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    report = compare_docx_reference(reference, ir)
    assert report["cer"] == 0.0
    assert report["wer"] == 0.0
